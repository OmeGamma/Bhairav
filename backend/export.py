from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
import io
from typing import Dict, Any

def generate_case_docx(case: Dict[str, Any]) -> io.BytesIO:
    doc = Document()
    
    # Title
    title = doc.add_heading(f"Case Report: {case.get('caseNumber', 'Unknown')}", level=0)
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    
    # Basic Info
    doc.add_heading('Basic Information', level=1)
    
    p = doc.add_paragraph()
    p.add_run('Title: ').bold = True
    p.add_run(f"{case.get('title', 'N/A')}\n")
    p.add_run('Crime Type: ').bold = True
    p.add_run(f"{case.get('crimeType', 'N/A')}\n")
    p.add_run('Status: ').bold = True
    p.add_run(f"{case.get('status', 'N/A')}\n")
    p.add_run('Priority: ').bold = True
    p.add_run(f"{case.get('priority', 'N/A')}\n")
    p.add_run('Officer: ').bold = True
    p.add_run(f"{case.get('investigatingOfficer', 'Unassigned')}\n")
    
    if case.get('location'):
        loc = case['location']
        loc_str = f"{loc.get('address', '')} {loc.get('city', '')} {loc.get('district', '')} {loc.get('state', '')}".strip()
        p.add_run('Location: ').bold = True
        p.add_run(f"{loc_str}\n")
    
    # Description
    doc.add_heading('Description', level=1)
    doc.add_paragraph(case.get('description') or case.get('incidentDetails') or 'No description provided.')
    
    if case.get('modusOperandi'):
        doc.add_heading('Modus Operandi', level=2)
        doc.add_paragraph(case['modusOperandi'])

    # Suspects & Persons
    suspects = case.get('suspects', [])
    victims = case.get('victims', [])
    persons = case.get('persons', [])
    
    if suspects or victims or persons:
        doc.add_heading('Involved Entities', level=1)
        
        if suspects:
            doc.add_heading('Suspects', level=2)
            for s in suspects:
                name = s.get('name', s) if isinstance(s, dict) else s
                doc.add_paragraph(f"- {name}")
                
        if victims:
            doc.add_heading('Victims', level=2)
            for v in victims:
                name = v.get('name', v) if isinstance(v, dict) else v
                doc.add_paragraph(f"- {name}")
                
        if persons:
            doc.add_heading('Other Persons of Interest', level=2)
            for p in persons:
                name = p.get('name', p) if isinstance(p, dict) else p
                doc.add_paragraph(f"- {name}")

    # Files / Evidence
    evidences = case.get('evidences', [])
    documents = case.get('documents', [])
    
    if evidences or documents:
        doc.add_heading('Evidence & Documents', level=1)
        for ev in evidences:
            label = ev.get("title") or ev.get("evidence_number") or ev.get("evidenceId") or "Evidence" if isinstance(ev, dict) else ev
            doc.add_paragraph(f"[EVIDENCE] {label}")
            
        for d in documents:
            label = d.get("title") or d.get("fileName") or d.get("documentId") or "Document" if isinstance(d, dict) else d
            doc.add_paragraph(f"[DOCUMENT] {label}")

    # Notes
    if case.get('notes'):
        doc.add_heading('Investigator Notes', level=1)
        doc.add_paragraph(case['notes'])

    # Save to BytesIO
    f = io.BytesIO()
    doc.save(f)
    f.seek(0)
    return f
